"""Document text extraction helpers."""

import csv
import datetime
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from src.utils.reading.ocr import ocr_pdf


def _read_csv_rows(csv_path: str) -> list[list[str]]:
    """Every row of a CSV, decoded and stripped. Never truncates."""

    last_error = None
    for encoding in ("utf-8", "latin-1"):
        try:
            with open(csv_path, newline="", encoding=encoding) as csv_file:
                sample = csv_file.read(4096)
                csv_file.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample)
                    # The sniffer can return a delimiter csv.reader then refuses
                    # ("bad delimiter value"), and that error is raised on read,
                    # outside the guard above. A one-character delimiter is the
                    # only thing csv.reader accepts, so check before trusting it.
                    if len(getattr(dialect, "delimiter", "") or "") != 1:
                        dialect = csv.excel
                except csv.Error:
                    dialect = csv.excel
                try:
                    return [
                        [cell.strip() for cell in row]
                        for row in csv.reader(csv_file, dialect)
                    ]
                except (csv.Error, ValueError):
                    csv_file.seek(0)
                    return [
                        [cell.strip() for cell in row]
                        for row in csv.reader(csv_file, csv.excel)
                    ]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Unable to decode CSV file: {last_error}")


def extract_csv_text(
    csv_path: str,
    max_rows: int = 500,
    max_chars: int = 100000,
) -> str:
    """Extract CSV content as tab-separated text for LLM analysis."""

    rows, duplicates = _dedup_rows(_read_csv_rows(csv_path))
    body = rows[:max_rows]
    if len(rows) > max_rows:
        body = body + [[f"... truncated after {max_rows} rows ..."]]
    if duplicates:
        body = [[f"... đã bỏ {duplicates} dòng trùng lặp ..."]] + body
    return "\n".join("\t".join(row) for row in body)[:max_chars]


def _format_number(value: float, number_format: str) -> str:
    """Render a numeric cell the way Excel displays it."""
    fmt = number_format or ""

    if "%" in fmt:
        scaled = value * 100
        text = f"{scaled:,.2f}".rstrip("0").rstrip(".")
        return _to_vietnamese_number(text or "0") + "%"

    # "#,##0" / "#,##0.00" ask Excel to group thousands.
    if "#,##" in fmt or "0,000" in fmt:
        decimals = 0
        if "." in fmt:
            decimals = len(fmt.split(".")[-1].split("_")[0].rstrip("%);\\ "))
        return _to_vietnamese_number(f"{value:,.{decimals}f}")

    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _to_vietnamese_number(us_formatted: str) -> str:
    """Swap US separators for Vietnamese ones ('.' thousands, ',' decimal)."""
    return us_formatted.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def _format_cell(value: object, number_format: str) -> str:
    """Render one cell value as faithful, LLM-readable text."""
    if value is None:
        return ""

    if isinstance(value, datetime.datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")

    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return _format_number(value, number_format)

    return str(value).strip()


def _sheet_grid(worksheet) -> list[list[str]]:
    """Return the sheet as rendered text cells, with merged values propagated."""
    grid = [
        [_format_cell(cell.value, cell.number_format) for cell in row]
        for row in worksheet.iter_rows()
    ]
    if not grid:
        return []

    used_width = max((len(row) for row in grid), default=0)
    for merged in worksheet.merged_cells.ranges:
        top, left = merged.min_row - 1, merged.min_col - 1
        if top >= len(grid) or left >= len(grid[top]):
            continue
        value = grid[top][left]
        if not value:
            continue
        if merged.min_col == 1 and merged.max_col >= used_width:
            continue
        for row_index in range(top, min(merged.max_row, len(grid))):
            for col_index in range(left, merged.max_col):
                if col_index < len(grid[row_index]):
                    grid[row_index][col_index] = value
    return grid


def _clean_grid(grid: list[list[str]]) -> tuple[list[list[str]], list[int]]:
    """Drop all-empty rows/columns; return the grid and its kept column indexes."""
    if not grid:
        return [], []

    width = max(len(row) for row in grid)
    padded = [row + [""] * (width - len(row)) for row in grid]
    kept_columns = [
        index for index in range(width) if any(row[index] for row in padded)
    ]
    cleaned = [
        [row[index] for index in kept_columns]
        for row in padded
        if any(row[index] for index in kept_columns)
    ]
    return cleaned, kept_columns


def _dedup_rows(grid: list[list[str]]) -> tuple[list[list[str]], int]:
    """Keep the first occurrence of each distinct row; also return how many went."""

    seen: set[tuple[str, ...]] = set()
    kept: list[list[str]] = []
    for row in grid:
        key = tuple(row)
        if key in seen:
            continue
        seen.add(key)
        kept.append(row)
    return kept, len(grid) - len(kept)


def _sheet_header(title: str, rows: int, columns: int, duplicates: int) -> str:
    """The "--- Sheet: … ---" line, naming duplicates when any were dropped."""

    dropped = f", đã bỏ {duplicates} dòng trùng lặp" if duplicates else ""
    return f"--- Sheet: {title} ({rows} dòng x {columns} cột{dropped}) ---"


def _grid_to_tsv(
    grid: list[list[str]],
    kept_columns: list[int],
    max_rows: int,
) -> str:
    """Render a cleaned grid as TSV led by its Excel column letters."""
    lines = ["\t".join(get_column_letter(index + 1) for index in kept_columns)]
    for row in grid[:max_rows]:
        # Right-trim trailing empties so short rows stop emitting runs of tabs.
        trimmed = list(row)
        while trimmed and not trimmed[-1]:
            trimmed.pop()
        lines.append("\t".join(trimmed))
    if len(grid) > max_rows:
        lines.append(f"... truncated after {max_rows} rows ...")
    return "\n".join(lines)


@dataclass(frozen=True)
class SheetText:
    """One sheet of a spreadsheet, already cleaned, de-duplicated and capped."""

    sheet_name: str
    content: str
    row_count: int
    column_count: int
    duplicates_dropped: int


def read_sheets(
    excel_path: str,
    max_rows_per_sheet: int = 500,
) -> list[SheetText]:
    """Every sheet of a spreadsheet, one record each."""

    suffix = Path(excel_path).suffix.lower()
    if suffix == ".csv":
        rows, duplicates = _dedup_rows(_read_csv_rows(excel_path))
        grid, kept_columns = _clean_grid(rows)
        if not grid:
            return [SheetText("", "", 0, 0, duplicates)]
        return [SheetText("", _grid_to_tsv(grid, kept_columns, max_rows_per_sheet),
                          len(grid), len(kept_columns), duplicates)]

    if suffix != ".xlsx":
        return _read_legacy_xls_sheets(excel_path, max_rows_per_sheet)

    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            grid, kept_columns = _clean_grid(_sheet_grid(worksheet))
            if not grid:
                sheets.append(SheetText(worksheet.title, "", 0, 0, 0))
                continue
            grid, duplicates = _dedup_rows(grid)
            sheets.append(SheetText(
                worksheet.title,
                _grid_to_tsv(grid, kept_columns, max_rows_per_sheet),
                len(grid), len(kept_columns), duplicates,
            ))
        return sheets
    finally:
        workbook.close()


def _read_legacy_xls_sheets(
    excel_path: str,
    max_rows_per_sheet: int,
) -> list[SheetText]:
    """A legacy .xls workbook, sheet by sheet (no number-format fidelity)."""

    workbook = pd.read_excel(
        excel_path, sheet_name=None, dtype=str, header=None, engine="xlrd",
    )
    sheets = []
    for sheet_name, dataframe in workbook.items():
        raw_grid = dataframe.fillna("").astype(str).values.tolist()
        grid, kept_columns = _clean_grid(raw_grid)
        if not grid:
            sheets.append(SheetText(sheet_name, "", 0, 0, 0))
            continue
        grid, duplicates = _dedup_rows(grid)
        sheets.append(SheetText(
            sheet_name, _grid_to_tsv(grid, kept_columns, max_rows_per_sheet),
            len(grid), len(kept_columns), duplicates,
        ))
    return sheets


def _join_sheet_blocks(blocks: list[str], max_chars: int) -> str:
    """Join sheet blocks, stopping with a visible marker at the char budget."""
    text = "\n\n".join(blocks)
    if len(text) <= max_chars:
        return text
    return (
        text[:max_chars].rstrip()
        + f"\n... truncated: workbook exceeds {max_chars} characters ..."
    )


def _sheets_to_text(sheets: list[SheetText], max_chars: int) -> str:
    """The one-string form every non-ledger caller reads."""

    blocks = [
        f"--- Sheet: {sheet.sheet_name} (trống) ---" if not sheet.row_count
        else _sheet_header(sheet.sheet_name, sheet.row_count,
                           sheet.column_count, sheet.duplicates_dropped)
        + f"\n{sheet.content}"
        for sheet in sheets
    ]
    return _join_sheet_blocks(blocks, max_chars)


def extract_excel_text(
    excel_path: str,
    max_rows_per_sheet: int = 500,
    max_chars: int = 100000,
) -> str:
    """Extract XLS/XLSX workbook content as LLM-readable text grouped by sheet."""

    return _sheets_to_text(read_sheets(excel_path, max_rows_per_sheet), max_chars)


def extract_xml_text(file_path: str, max_chars: int = 120_000) -> str:
    """Flatten an XML file to readable text."""

    import xml.etree.ElementTree as ET

    raw = Path(file_path).read_text(encoding="utf-8-sig", errors="replace")
    start = raw.find("<?xml")
    if start > 0:
        raw = raw[start:]
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        return f"[XML không đọc được: {exc}]\n\n{raw[:max_chars]}"

    lines: list[str] = []

    def walk(node, path: str) -> None:
        tag = node.tag.split("}")[-1]
        here = f"{path}/{tag}" if path else tag
        text = (node.text or "").strip()
        if text:
            lines.append(f"{here}: {text}")
        for child in node:
            walk(child, here)

    walk(root, "")
    return "\n".join(lines)[:max_chars]


def extract_document_text(
    file_path: str,
    ocr_timeout_seconds: float | None = None,
) -> str:
    """Extract text from supported uploaded documents."""
    
    spreadsheet_extensions = {".csv", ".xls", ".xlsx"}

    extension = Path(file_path).suffix.lower()
    if extension == ".pdf":
        return ocr_pdf(file_path, timeout_seconds=ocr_timeout_seconds)
    if extension == ".csv":
        return extract_csv_text(file_path)
    if extension in spreadsheet_extensions - {".csv"}:
        return extract_excel_text(file_path)
    if extension == ".xml":
        return extract_xml_text(file_path)
    if extension == ".docx":
        return extract_docx_text(file_path)
    if extension in {".txt", ".md"}:
        return Path(file_path).read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file extension: {extension}")


def extract_docx_text(file_path: str) -> str:
    """Đoạn văn và bảng của một file Word, theo thứ tự xuất hiện."""

    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(file_path)
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text.strip()
            if text:
                parts.append(text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, document).rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    return "\n".join(parts)
